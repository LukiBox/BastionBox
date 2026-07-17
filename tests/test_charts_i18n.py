"""Charts & photos in AI documents + the Polish language layer.

Two features, both end-to-end:
* write_document renders ```chart JSON fences as real charts (vector in PDF,
  PNG in Word) and embeds ![caption](path) photos resolved through the jail —
  workspace or read-only library — with clear errors for bad references.
* The i18n catalog has full EN/PL parity, and the Settings language switcher
  retranslates the live window and persists the choice to the encrypted store.
"""
from __future__ import annotations

import base64
import io
import re

import pytest

from bastion.core.agent.permissions import Decision, PolicyBroker
from bastion.core.security.audit import AuditLog
from bastion.core.security.jail import PathJail, Permission
from bastion.core.tools.base import ToolContext

docx = pytest.importorskip("docx")
fitz = pytest.importorskip("fitz")

# A valid 1×1 red PNG for photo-embedding tests.
_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGP4z8DwHwAFAAH/"
    "q842iQAAAABJRU5ErkJggg==")

_CHART_MD = """# Environmental Test Report

Measured power draw during the thermal cycle.

```chart
{"type": "bar", "title": "Power draw by mode", "labels": ["Idle", "Cruise", "Peak"],
 "series": [{"name": "Unit A", "values": [4.2, 18, 45]},
            {"name": "Unit B", "values": [3.9, 17, 44]}]}
```

![Test rig on the shaker table](photos/rig.png)

Conclusion paragraph.
"""


@pytest.fixture()
def ctx(tmp_path):
    ws_dir = tmp_path / "ws"; (ws_dir / "photos").mkdir(parents=True)
    lib_dir = tmp_path / "lib"; lib_dir.mkdir()
    (ws_dir / "photos" / "rig.png").write_bytes(_PNG)
    jail = PathJail()
    ws = jail.mount(ws_dir, Permission.ASK)
    lib = jail.mount(lib_dir, Permission.READ_ONLY, label="norms")
    broker = PolicyBroker(ask_write=lambda w, d: Decision(True, "ok"))
    return (ToolContext(jail=jail, workspace=ws, broker=broker,
                        audit=AuditLog(tmp_path / "a.jsonl"), library=lib),
            ws_dir, lib_dir)


# -- chart engine -------------------------------------------------------------
def test_chart_png_all_types():
    from bastion.core.docs.charts import chart_png
    spec = {"labels": ["A", "B"], "series": [{"name": "S", "values": [1, 2]}]}
    for kind in ("bar", "line", "pie"):
        png = chart_png(dict(spec, type=kind, title=kind.upper()))
        assert png[:4] == b"\x89PNG", kind
        assert len(png) > 2_000, kind          # a real render, not a stub


def test_chart_spec_is_tolerant_of_messy_model_output():
    from bastion.core.docs.charts import normalize_spec
    # Flat list, string numbers with a European decimal comma, short labels.
    spec = normalize_spec({"labels": ["a"], "series": [1, "2,5", 3]})
    assert spec["series"][0]["values"] == [1.0, 2.5, 3.0]
    assert len(spec["labels"]) == 3            # padded to match the data
    # Mapping form {name: values}.
    spec = normalize_spec({"series": {"X": [1], "Y": [2]}, "labels": ["q"]})
    assert [s["name"] for s in spec["series"]] == ["X", "Y"]


def test_chart_spec_errors_name_the_problem():
    from bastion.core.docs.charts import ChartError, normalize_spec
    with pytest.raises(ChartError, match="valid JSON"):
        normalize_spec("{nope")
    with pytest.raises(ChartError, match="at least one series"):
        normalize_spec({"series": []})
    with pytest.raises(ChartError, match="donut"):
        normalize_spec({"type": "donut", "series": [1]})
    with pytest.raises(ChartError, match="not a number"):
        normalize_spec({"series": [{"values": ["fast"]}]})


# -- write_document embeds charts + photos -------------------------------------
def test_docx_report_embeds_chart_and_photo(ctx):
    from bastion.core.tools.office_tools import WriteDocument
    c, ws_dir, _ = ctx
    r = WriteDocument().run(c, {"path": "report.docx", "title": "Thermal",
                                "content": _CHART_MD})
    assert r.ok, r.observation
    d = docx.Document(io.BytesIO((ws_dir / "report.docx").read_bytes()))
    assert len(d.inline_shapes) == 2          # rendered chart + the photo
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Test rig on the shaker table" in text   # caption survived
    assert "```" not in text                        # fence never leaks as text


def test_pdf_report_embeds_vector_chart_and_photo(ctx):
    from bastion.core.tools.office_tools import WriteDocument
    c, ws_dir, _ = ctx
    r = WriteDocument().run(c, {"path": "report.pdf", "title": "Thermal",
                                "content": _CHART_MD})
    assert r.ok, r.observation
    with fitz.open(ws_dir / "report.pdf") as pdf:
        page = pdf[0]
        assert page.get_images(), "photo missing from the PDF"
        assert page.get_drawings(), "vector chart missing from the PDF"
        assert "Power draw by mode" in page.get_text()   # chart title is text


def test_photo_resolves_from_readonly_library(ctx):
    from bastion.core.tools.office_tools import WriteDocument
    c, ws_dir, lib_dir = ctx
    (lib_dir / "norm_fig.png").write_bytes(_PNG)
    r = WriteDocument().run(c, {
        "path": "cited.docx", "title": "T",
        "content": "![Figure 3 from the norm](norm_fig.png)"})
    assert r.ok, r.observation
    d = docx.Document(io.BytesIO((ws_dir / "cited.docx").read_bytes()))
    assert len(d.inline_shapes) == 1


def test_missing_photo_is_a_clear_error_and_nothing_is_written(ctx):
    from bastion.core.tools.office_tools import WriteDocument
    c, ws_dir, _ = ctx
    r = WriteDocument().run(c, {"path": "bad.docx", "title": "T",
                                "content": "![x](photos/nope.png)"})
    assert not r.ok
    assert "photos/nope.png" in r.observation
    assert not (ws_dir / "bad.docx").exists()


def test_non_image_reference_is_refused(ctx):
    from bastion.core.tools.office_tools import WriteDocument
    c, ws_dir, _ = ctx
    (ws_dir / "notes.txt").write_text("hi", encoding="utf-8")
    r = WriteDocument().run(c, {"path": "bad.docx", "title": "T",
                                "content": "![x](notes.txt)"})
    assert not r.ok and "unsupported" in r.observation


def test_bad_chart_json_is_a_clear_error(ctx):
    from bastion.core.tools.office_tools import WriteDocument
    c, ws_dir, _ = ctx
    r = WriteDocument().run(c, {"path": "bad.docx", "title": "T",
                                "content": "```chart\n{oops\n```"})
    assert not r.ok and "chart block" in r.observation
    assert not (ws_dir / "bad.docx").exists()


# -- write_spreadsheet native chart --------------------------------------------
def test_xlsx_gets_a_native_editable_chart(ctx):
    import openpyxl
    from bastion.core.tools.office_tools import WriteSpreadsheet
    c, ws_dir, _ = ctx
    r = WriteSpreadsheet().run(c, {
        "path": "data.xlsx", "sheet": "Results",
        "data": "Mode,Unit A,Unit B\nIdle,4.2,3.9\nPeak,45,44",
        "chart": '{"type": "line", "title": "Draw"}'})
    assert r.ok, r.observation
    wb = openpyxl.load_workbook(io.BytesIO((ws_dir / "data.xlsx").read_bytes()))
    assert len(wb["Results"]._charts) == 1
    # And the data itself is intact next to the chart.
    assert wb["Results"]["B3"].value == 45


def test_xlsx_bad_chart_type_is_refused(ctx):
    from bastion.core.tools.office_tools import WriteSpreadsheet
    c, ws_dir, _ = ctx
    r = WriteSpreadsheet().run(c, {"path": "d.xlsx", "data": "a,b\n1,2",
                                   "chart": '{"type": "scatter"}'})
    assert not r.ok and "scatter" in r.observation
    assert not (ws_dir / "d.xlsx").exists()


# -- i18n catalog ---------------------------------------------------------------
_PLACEHOLDER = re.compile(r"{(\w+)}")


def test_catalog_full_parity_between_english_and_polish():
    from bastion.core.i18n import TRANSLATIONS
    en, pl = TRANSLATIONS["en"], TRANSLATIONS["pl"]
    assert set(en) == set(pl), (
        "keys differ: " + str(set(en) ^ set(pl)))
    for key, text in en.items():
        assert text.strip() and pl[key].strip(), key
        # Format placeholders must match, or a switch breaks .format() calls.
        assert set(_PLACEHOLDER.findall(text)) == \
            set(_PLACEHOLDER.findall(pl[key])), key


def test_translator_fallbacks():
    from bastion.core.i18n import Translator
    tr = Translator("xx")                     # unknown language → English
    assert tr.language == "en"
    tr = Translator("pl")
    assert tr.t("totally.bogus.key") == "totally.bogus.key"
    assert tr.t("nav.chat") == "Czat"


def test_system_prompt_documents_charts_and_photos():
    from bastion.core.agent.schemas import build_system_prompt, default_toolbox
    p = build_system_prompt(default_toolbox(), "ws", "ask")
    assert "```chart" in p and "![caption](photos/rig.png)" in p
