"""Workflow readiness — the four eigent-class benchmark tasks, in miniature.

Each test drives the *real* tools (and where it matters, the real agent loop)
through a deterministic script shaped like the user's benchmark prompts:
CSV → HTML report with a chart, duplicate scan, physics calculations for a lab
PDF, and invoices → Excel BOM. Plus safety tests proving `calculate` is math
and nothing else.
"""
from __future__ import annotations

import json

import pytest

from bastion.core.agent.loop import AgentLoop, EventKind
from bastion.core.agent.permissions import AutoApproveBroker
from bastion.core.llm.engine import FakeEngine
from bastion.core.security.audit import AuditLog
from bastion.core.security.jail import PathJail, Permission
from bastion.core.tools.base import ToolContext
from bastion.core.tools.compute import Calculate, FindDuplicates, evaluate, _CalcError


def _call(tool, **args):
    return json.dumps({"tool": tool, "args": args})


def _final(content):
    return json.dumps({"tool": "final", "args": {"content": content}})


@pytest.fixture()
def env(tmp_path):
    ws_dir = tmp_path / "ws"
    ws_dir.mkdir()
    jail = PathJail()
    ws = jail.mount(ws_dir, Permission.ASK, label="ws")
    audit = AuditLog(tmp_path / "audit.jsonl")
    return jail, ws, audit, ws_dir


def _ctx(env):
    jail, ws, audit, _ = env
    return ToolContext(jail=jail, workspace=ws, broker=AutoApproveBroker(),
                       audit=audit)


# -- calculate: the lab-report physics, exactly -------------------------------
def test_calculate_lab_physics(env):
    # The user's benchmark: mmHg→Pa, ρ from p and T, Reynolds number.
    r = Calculate().run(_ctx(env), {"expression":
        "p = 751 * 133.322; T = 26.5 + 273.15; rho = p / (287.05 * T); "
        "v = 14.2; L = 0.05; Re = v * L / (16.1 * 10**-6); Re"})
    assert r.ok
    vals = dict(line.split(" = ") for line in r.observation.splitlines())
    assert abs(float(vals["p"]) - 100124.822) < 0.01
    assert abs(float(vals["rho"]) - 100124.822 / (287.05 * 299.65)) < 1e-9
    assert abs(float(vals["Re"]) - 14.2 * 0.05 / 16.1e-6) < 1e-3
    assert r.meta["result"] == pytest.approx(44099.3788, abs=1e-3)


def test_calculate_helpful_errors_and_functions(env):
    c = Calculate()
    assert "unknown variable" in c.run(_ctx(env), {"expression": "x + 1"}).observation
    assert "division by zero" in c.run(_ctx(env), {"expression": "1/0"}).observation
    ok = c.run(_ctx(env), {"expression": "sqrt(2) * cos(0) + log10(100)"})
    assert ok.ok and ok.meta["result"] == pytest.approx(2**0.5 + 2)


def test_calculate_is_math_not_python():
    for evil in ("__import__('os')", "open('x')", "(1).__class__",
                 "[i for i in range(9)]", "'a' * 9", "lambda: 1",
                 "10**10**10", "9**9999"):
        with pytest.raises(_CalcError):
            evaluate(evil)


# -- workflow 2: duplicate scan ------------------------------------------------
def test_find_duplicates_groups_exact_near_and_size(env):
    ws_dir = env[3]
    docs = ws_dir / "mydocs"
    (docs / "sub").mkdir(parents=True)
    (docs / "report.txt").write_bytes(b"Q2 spend was 42000 EUR total.")
    (docs / "sub" / "copy_of_report.bak").write_bytes(
        b"Q2 spend was 42000 EUR total.")            # exact dupe, other name/ext
    (docs / "report_v2.txt").write_bytes(
        b"Q2  spend was\n42000 EUR    total.")        # near-dupe (whitespace)
    (docs / "aaaa.bin").write_bytes(b"\x00\x01\x02\x03")
    (docs / "bbbb.bin").write_bytes(b"\x04\x05\x06\x07")  # same size, diff bytes
    (docs / "unique.txt").write_bytes(b"nothing like the others at all")

    r = FindDuplicates().run(_ctx(env), {"path": "mydocs"})
    assert r.ok
    assert r.meta["exact_groups"] == 1 and r.meta["near_groups"] == 1
    assert "EXACT DUPLICATES" in r.observation
    assert "report.txt" in r.observation and "copy_of_report.bak" in r.observation
    assert "NEAR-DUPLICATES" in r.observation and "report_v2.txt" in r.observation
    assert "SAME SIZE" in r.observation and "aaaa.bin" in r.observation


def test_find_duplicates_clean_folder(env):
    ws_dir = env[3]
    (ws_dir / "one.txt").write_text("alpha", encoding="utf-8")
    (ws_dir / "two.txt").write_text("beta together now", encoding="utf-8")
    r = FindDuplicates().run(_ctx(env), {"path": "."})
    assert r.ok and "no duplicates found" in r.observation


# -- workflow 1: CSV → HTML report with chart -----------------------------------
def test_csv_to_html_report_with_chart(env):
    ws_dir = env[3]
    (ws_dir / "bank_transactions.csv").write_text(
        "date,category,amount\n2026-04-03,Cloud,1200\n2026-05-11,Cloud,1350\n"
        "2026-06-20,Office,800\n", encoding="utf-8")
    content = (
        "## Q2 Spend\n\nTotal spend was **3350 EUR**.\n\n"
        "category | total\nCloud | 2550\nOffice | 800\n\n"
        "```chart\n"
        '{"type": "bar", "title": "Q2 spend by category",'
        ' "labels": ["Cloud", "Office"],'
        ' "series": [{"name": "EUR", "values": [2550, 800]}]}\n'
        "```\n")
    script = [
        _call("read_document", path="bank_transactions.csv"),
        _call("calculate", expression="cloud = 1200 + 1350; total = cloud + 800; total"),
        _call("write_document", path="q2_report.html",
              title="Q2 Financial Statement", content=content),
        _final("Wrote q2_report.html: 3350 EUR total, chart included."),
    ]
    events = list(AgentLoop(FakeEngine(script), _ctx(env)).run(
        "prepare a Q2 statement from bank_transactions.csv as an html report"))
    assert events[-1].kind is EventKind.FINAL
    obs = [e.text for e in events if e.kind is EventKind.OBSERVATION]
    assert any("total = 3350" in o for o in obs)        # exact math happened
    html = (ws_dir / "q2_report.html").read_text("utf-8")
    assert "<!DOCTYPE html>" in html
    assert "Q2 Financial Statement" in html
    assert "data:image/png;base64," in html             # chart embedded
    assert "<strong>3350 EUR</strong>" in html           # emphasis rendered
    assert "<th>category</th>" in html                   # table rendered


def test_html_escapes_hostile_text(env):
    from bastion.core.docs.writer import build_html
    page = build_html([{"type": "paragraph",
                        "text": "<script>alert(1)</script> & co"}],
                      title="T<i>tle").decode("utf-8")
    assert "<script>" not in page
    assert "&lt;script&gt;" in page and "&amp; co" in page


# -- workflow 4: invoices folder → Excel BOM ------------------------------------
def test_invoices_to_excel_bom(env):
    ws_dir = env[3]
    inv = ws_dir / "invoices"
    inv.mkdir()
    (inv / "inv_001.txt").write_text(
        "Invoice 001\nM4 bolts x200 @ 0.12\nAlu plate x4 @ 18.50", "utf-8")
    (inv / "inv_002.txt").write_text(
        "Invoice 002\nBearings x8 @ 6.40\nM4 bolts x100 @ 0.12", "utf-8")
    script = [
        _call("list_dir", path="invoices"),
        _call("read_document", path="invoices/inv_001.txt"),
        _call("append_note", name="bom",
              content="M4 bolts,200,0.12\nAlu plate,4,18.50"),
        _call("read_document", path="invoices/inv_002.txt"),
        _call("append_note", name="bom", content="Bearings,8,6.40\nM4 bolts,100,0.12"),
        _call("calculate", expression="bolts = 200 + 100; bolts"),
        _call("write_spreadsheet", path="bom.xlsx",
              data=("item,qty,unit_price\nM4 bolts,300,0.12\n"
                    "Alu plate,4,18.50\nBearings,8,6.40"),
              sheet="BOM"),
        _final("bom.xlsx written: 3 items consolidated from 2 invoices."),
    ]
    events = list(AgentLoop(FakeEngine(script), _ctx(env)).run(
        "create a bill of materials from the invoices folder"))
    assert events[-1].kind is EventKind.FINAL
    assert (ws_dir / "bom.xlsx").exists()
    import openpyxl, io
    wb = openpyxl.load_workbook(io.BytesIO((ws_dir / "bom.xlsx").read_bytes()))
    rows = list(wb["BOM"].values)
    assert rows[0] == ("item", "qty", "unit_price")
    assert ("M4 bolts", 300, 0.12) in rows              # consolidated quantity


# -- the 20-minute-hang class of bugs -------------------------------------------
def test_read_file_refuses_documents_and_binaries(env):
    """A .docx through read_file used to flood the context with 200KB of binary
    garbage — one step then took many minutes of CPU prefill. Now it redirects."""
    ws_dir = env[3]
    (ws_dir / "report.docx").write_bytes(b"PK\x03\x04 fake zip bytes" * 100)
    (ws_dir / "blob.dat").write_bytes(b"\x00\x01\x02" * 200)
    from bastion.core.tools.fs_tools import ReadFile
    r = ReadFile().run(_ctx(env), {"path": "report.docx"})
    assert not r.ok and "read_document" in r.observation
    r = ReadFile().run(_ctx(env), {"path": "blob.dat"})
    assert not r.ok and "binary" in r.observation


def test_read_file_respects_context_budget(env):
    ws_dir = env[3]
    (ws_dir / "big.log").write_text("x" * 50_000, encoding="utf-8")
    ctx = _ctx(env)
    ctx.read_char_cap = 10_000
    from bastion.core.tools.fs_tools import ReadFile
    r = ReadFile().run(ctx, {"path": "big.log"})
    assert r.ok and len(r.observation) < 11_000 and "truncated" in r.observation


def test_not_found_path_hints_at_workspace_root(env):
    from bastion.core.tools.fs_tools import ListDir
    r = ListDir().run(_ctx(env), {"path": "test"})   # the user's exact miss
    assert not r.ok and "'.'" in r.observation


def test_engine_failure_ends_gracefully_not_silently(env):
    class DyingEngine(FakeEngine):
        def stream(self, messages, config):
            raise ConnectionError("ollama server went away")
    events = list(AgentLoop(DyingEngine([]), _ctx(env)).run("do something"))
    errs = [e for e in events if e.kind is EventKind.ERROR]
    assert any("backend failed" in e.text and "ollama server went away" in e.text
               for e in errs)
    assert events[-1].kind is EventKind.FINAL           # still a usable ending
    assert events[-1].meta.get("partial") is True


def test_progress_events_stream_during_generation(env):
    long_final = _final("A" * 900)                       # forces >4 chunks
    events = list(AgentLoop(FakeEngine([long_final]), _ctx(env)).run("hi"))
    progress = [e for e in events if e.kind is EventKind.PROGRESS]
    assert progress and progress[0].meta["chars"] == 0   # "reading context"
    assert any(p.meta["chars"] > 0 for p in progress)    # then generation pings


def test_ollama_backend_carries_context():
    from bastion.core.llm.ollama_backend import OllamaBackend
    b = OllamaBackend("qwen3:30b", n_ctx=16384)
    assert b.n_ctx == 16384   # requested on every generate via options.num_ctx


def test_write_document_rejects_plan_checklist_as_content(env):
    """The 'plan pasted as document' failure from a real session: a short body
    of unchecked checkboxes is the notepad leaking, not a report."""
    from bastion.core.tools.office_tools import WriteDocument
    r = WriteDocument().run(_ctx(env), {
        "path": "summary.html", "title": "Summary",
        "content": "## Summary of Findings\n- [ ] create summary document\n"
                   "- [ ] add content"})
    assert not r.ok and "checklist" in r.observation
    # A real report with one incidental checkbox line still writes fine.
    r2 = WriteDocument().run(_ctx(env), {
        "path": "report.html", "title": "EMC Report",
        "content": "## Findings\n\nNFR_EMC_02 passed at 45 W peak draw. "
                   + "Details follow. " * 40 + "\n- [ ] pending retest item\n"
                   "- [ ] spare fixture check"})
    assert r2.ok


def test_loop_history_reaches_the_prompt(env):
    """history passed to run() must be visible to the model — the UI-side
    regression is covered in test_ui_smoke; this pins the loop contract."""
    from bastion.core.llm.engine import Message, Role

    class Capture(FakeEngine):
        prompts = []
        def stream(self, messages, config):
            Capture.prompts.append(list(messages))
            yield from super().stream(messages, config)

    history = [Message(Role.USER, "what is NFR_EMC_02?"),
               Message(Role.ASSISTANT, "A conducted-emissions test at 45 W.")]
    events = list(AgentLoop(Capture([_final("done")]), _ctx(env)).run(
        "summarize what you learned", history=history))
    assert events[-1].kind is EventKind.FINAL
    joined = "\n".join(m.content for m in Capture.prompts[0])
    assert "NFR_EMC_02" in joined and "conducted-emissions" in joined
